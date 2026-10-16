#!/usr/bin/env python3
"""Update CO validation workbooks with project dates/status columns and MW checks.

By default this script runs in dry-run mode and prints the changes it would make.
Use --apply to export updated workbook copies to the output directory. Source
workbooks are not modified.
"""

from __future__ import annotations

import argparse
import copy
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Iterable

import openpyxl
from openpyxl.comments import Comment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.styles import Alignment, Font, PatternFill


DEFAULT_INPUT_DIR = Path("01_Input/00_CO Validation - raw")
DEFAULT_OUTPUT_DIR = Path("01_Input/00_CO Validation")

PROJECT_INFO_HEADERS_AFTER_BUDGET = ["Start Year", "End Year", "Status", "VF or Non-VF"]
PROJECT_INFO_HEADERS = [
    "Project ID",
    "Title",
    "Link",
    "Budget",
    "Budget (USD)",
    *PROJECT_INFO_HEADERS_AFTER_BUDGET,
]
DATE_STATUS_HEADERS = ["Start Year", "End Year", "Status"]
STATUS_OPTIONS = [
    "Pending",
    "Approved",
    "Active",
    "Operationally Closed",
    "Financially Closed",
    "Suspended",
    "Cancelled",
]
STATUS_ALIASES = {
    "Finalised": "Financially Closed",
    "Finalized": "Financially Closed",
}
VF_ALIASES = {
    "VF": "VF",
    "Vertical Funded": "VF",
    "Vertically Funded": "VF",
    "Non-VF": "Non-VF",
    "non-VF": "Non-VF",
    "NON-VF": "Non-VF",
    "Non VF": "Non-VF",
    "NonVF": "Non-VF",
    "NFV": "Non-VF",
}
BENEFICIARY_CATEGORIES_SHEET = "Beneficiary Categories"
AUTO_NOTE = (
    "[AUTO-ADDED] Project reports access/productive-use beneficiary numbers "
    "but no MW value. Please add installed capacity (MW) or confirm not applicable."
)
NEEDS_MW_STATUS = "Needs MW value"

ACTION_NEEDED_FILL = PatternFill(fill_type="solid", fgColor="FFFF00")

BENEFICIARY_CATEGORIES = [
    (
        "Electricity Access",
        "Number of people, disaggregated by sex, who gain access to clean electricity (direct access to electricity, lighting, cooling, etc.).",
    ),
    ("Energy (MW added)", "Installed renewable or clean-energy power capacity added, in MW."),
    ("Energy Storage (MW)", "Installed energy storage power capacity added, in MW."),
    (
        "Clean Cooking",
        "Number of people who gain access to clean cooking (direct access to clean cook stoves, clean fuels, biomass, etc.).",
    ),
    (
        "Clean Heating",
        "Number of people who gain access to clean heating (direct access to clean electric heaters, clean fuels, etc.).",
    ),
    (
        "Agriculture and Food System",
        "Number of people who gain access to clean, affordable, and sustainable energy for agricultural and food-system activities.",
    ),
    (
        "Health Services",
        "Number of people who gain access to health care facilities through clean energy systems.",
    ),
    (
        "Water Services",
        "Number of people who gain access to water services through clean energy systems.",
    ),
    (
        "Education Services",
        "Number of people who gain access to education services through clean energy systems, including students and teachers.",
    ),
    (
        "Transportation and E-mobility Services",
        "Number of people served or energy saved from transport and e-mobility interventions, such as electric vehicles or charging stations.",
    ),
    (
        "Energy Efficiency Services",
        "Energy saved from energy-efficiency interventions, such as building or industrial efficiency.",
    ),
    (
        "Energy Infrastructure Services",
        "Number of people benefiting from energy infrastructure interventions, such as streetlights, transmission, and distribution lines.",
    ),
    ("Other Energy Services", "Access to other energy services not captured by the categories above."),
    ("Financing support", "Support for the development of the energy market ecosystem."),
    (
        "Medium Enterprises",
        "Support for the development of medium enterprises in the energy transition market.",
    ),
    (
        "Small Enterprises",
        "Support for the development of small enterprises in the energy transition market.",
    ),
    (
        "Capacity Building Training",
        "Number of people benefiting from training on energy business, energy finance, or related implementation capacity.",
    ),
    (
        "Entrepreneurship Training",
        "Training for enterprises on the energy business, assumed to result in direct beneficiaries.",
    ),
    (
        "Market Intervention",
        "Volume of investment leveraged or barriers removed for clean, reliable, affordable, and sustainable energy solutions.",
    ),
    ("Campaign Participant", "Individuals who participate in advocacy or awareness campaigns on energy."),
    (
        "Policy or Regulatory Framework",
        "Policies, regulations, strategies, plans, or frameworks developed, approved, implemented, or strengthened.",
    ),
    ("GHG Emissions Reduction", "Greenhouse gas emissions reduced or mitigated, in tCO2e."),
    ("Other", "Energy-related output that does not fit the standard categories above."),
]

MW_CATEGORY_RE = re.compile(r"\bmw\b", re.IGNORECASE)
NUMBER_RE = re.compile(r"[-+]?\d[\d,]*(?:\.\d+)?")

INCLUDE_CATEGORY_TERMS = {
    "access",
    "clean cooking",
    "clean heating",
    "agriculture",
    "food system",
    "health services",
    "education services",
    "water services",
    "transportation",
    "transport services",
    "e-mobility",
    "small enterprises",
    "medium enterprises",
    "entrepreneurship",
    "income generating",
    "livelihood",
}

EXCLUDE_CATEGORY_TERMS = {
    "capacity",
    "campaign",
    "efficiency",
    "financing",
    "ghg",
    "market",
    "policy",
    "regulatory",
    "training",
}


@dataclass
class SheetChange:
    workbook: Path
    sheet: str
    columns_added: list[str] = field(default_factory=list)
    placeholder_rows_added: int = 0
    existing_mw_rows_marked: int = 0
    project_rows_highlighted: int = 0
    output_rows_project_info_cleared: int = 0
    status_values_normalized: int = 0
    vf_values_normalized: int = 0
    status_validation_added: bool = False
    beneficiary_categories_updated: bool = False

    @property
    def changed(self) -> bool:
        return bool(
            self.columns_added
            or self.placeholder_rows_added
            or self.existing_mw_rows_marked
            or self.project_rows_highlighted
            or self.output_rows_project_info_cleared
            or self.status_values_normalized
            or self.vf_values_normalized
            or self.status_validation_added
            or self.beneficiary_categories_updated
        )


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description=(
            "Add Start Year, End Year, and Status columns after Budget in CO "
            "validation workbooks, and flag projects that report access/productive "
            "use beneficiary targets without a MW value."
        )
    )
    parser.add_argument(
        "--input-dir",
        type=Path,
        default=DEFAULT_INPUT_DIR,
        help=f"Directory of CO validation .xlsx files. Default: {DEFAULT_INPUT_DIR}",
    )
    parser.add_argument(
        "--output-dir",
        type=Path,
        default=DEFAULT_OUTPUT_DIR,
        help=f"Directory for updated workbook copies. Default: {DEFAULT_OUTPUT_DIR}",
    )
    parser.add_argument(
        "--apply",
        action="store_true",
        help=(
            "Export updated workbook copies to --output-dir. Without this flag "
            "the script only reports changes."
        ),
    )
    parser.add_argument(
        "--sheet-name",
        default="Projects (3)",
        help="Worksheet to update by default. Default: 'Projects (3)'.",
    )
    parser.add_argument(
        "--all-project-sheets",
        action="store_true",
        help=(
            "Update every sheet with validation-table headers instead of only "
            "--sheet-name."
        ),
    )
    return parser.parse_args()


def normalize(value: object) -> str:
    return str(value).strip() if value is not None else ""


def normalized_header_map(ws, header_row: int) -> dict[str, int]:
    headers: dict[str, int] = {}
    for col in range(1, ws.max_column + 1):
        header = normalize(ws.cell(header_row, col).value)
        if header and header not in headers:
            headers[header] = col
    return headers


def budget_col_from_headers(headers: dict[str, int]) -> int | None:
    return headers.get("Budget") or headers.get("Budget (USD)")


def normalize_status_value(value: object) -> object:
    text = normalize(value)
    return STATUS_ALIASES.get(text, value)


def normalize_vf_value(value: object) -> object:
    text = normalize(value).replace("\u00a0", " ")
    text = re.sub(r"\s+", " ", text)
    return VF_ALIASES.get(text, value)


def find_header_row(ws) -> int | None:
    for row in range(1, min(ws.max_row, 5) + 1):
        headers = set(normalize(cell.value) for cell in ws[row])
        if {"Project ID", "Beneficiary Category"}.issubset(headers) and (
            "Budget" in headers or "Budget (USD)" in headers
        ):
            return row
    return None


def copy_column_style(ws, source_col: int, target_col: int) -> None:
    ws.column_dimensions[
        openpyxl.utils.get_column_letter(target_col)
    ].width = max(
        ws.column_dimensions[
            openpyxl.utils.get_column_letter(source_col)
        ].width
        or 0,
        12,
    )
    for row in range(1, ws.max_row + 1):
        source = ws.cell(row, source_col)
        target = ws.cell(row, target_col)
        if source.has_style:
            target._style = copy.copy(source._style)
        if source.number_format:
            target.number_format = source.number_format
        if source.alignment:
            target.alignment = copy.copy(source.alignment)


def unmerge_sheet_cells(ws) -> int:
    merged_ranges = list(ws.merged_cells.ranges)
    for merged_range in merged_ranges:
        ws.unmerge_cells(str(merged_range))
    return len(merged_ranges)


def ensure_project_info_columns(ws, header_row: int) -> list[str]:
    headers = normalized_header_map(ws, header_row)
    budget_col = budget_col_from_headers(headers)
    if budget_col is None:
        return []

    desired_cols = list(
        range(
            budget_col + 1,
            budget_col + 1 + len(PROJECT_INFO_HEADERS_AFTER_BUDGET),
        )
    )
    current = [
        normalize(ws.cell(header_row, col).value)
        for col in desired_cols
        if col <= ws.max_column
    ]
    if current == PROJECT_INFO_HEADERS_AFTER_BUDGET:
        return []

    values_by_header: dict[str, list[object]] = {}
    original_headers = normalized_header_map(ws, header_row)
    existing_cols = [
        original_headers[header]
        for header in PROJECT_INFO_HEADERS_AFTER_BUDGET
        if header in original_headers
    ]
    for header in PROJECT_INFO_HEADERS_AFTER_BUDGET:
        if header in original_headers:
            col = original_headers[header]
            values_by_header[header] = [
                ws.cell(row, col).value for row in range(header_row + 1, ws.max_row + 1)
            ]

    for col in sorted(existing_cols, reverse=True):
        ws.delete_cols(col)

    headers = normalized_header_map(ws, header_row)
    budget_col = budget_col_from_headers(headers)
    if budget_col is None:
        return []
    insert_at = budget_col + 1
    ws.insert_cols(insert_at, len(PROJECT_INFO_HEADERS_AFTER_BUDGET))

    for offset, header in enumerate(PROJECT_INFO_HEADERS_AFTER_BUDGET):
        col = insert_at + offset
        copy_column_style(ws, budget_col, col)
        ws.cell(header_row, col).value = header
        for row_offset, value in enumerate(values_by_header.get(header, []), start=1):
            row = header_row + row_offset
            if row <= ws.max_row:
                ws.cell(row, col).value = (
                    normalize_status_value(value) if header == "Status"
                    else normalize_vf_value(value) if header == "VF or Non-VF"
                    else value
                )

    return PROJECT_INFO_HEADERS_AFTER_BUDGET


def ensure_status_data_validation(ws, header_row: int) -> bool:
    headers = normalized_header_map(ws, header_row)
    status_col = headers.get("Status")
    if status_col is None:
        return False

    status_letter = openpyxl.utils.get_column_letter(status_col)
    first_data_row = header_row + 1
    last_data_row = max(ws.max_row, header_row + 2000)
    formula = f'"{",".join(STATUS_OPTIONS)}"'
    validation = DataValidation(
        type="list",
        formula1=formula,
        allow_blank=True,
        showDropDown=False,
    )
    validation.error = "Select a status from the dropdown list."
    validation.errorTitle = "Invalid Status"
    validation.prompt = "Choose a project status."
    validation.promptTitle = "Status"
    validation.add(f"{status_letter}{first_data_row}:{status_letter}{last_data_row}")
    ws.add_data_validation(validation)
    return True


def normalize_status_cells(ws, header_row: int) -> int:
    headers = normalized_header_map(ws, header_row)
    status_col = headers.get("Status")
    if status_col is None:
        return 0

    normalized = 0
    for row in range(header_row + 1, ws.max_row + 1):
        cell = ws.cell(row, status_col)
        cleaned = normalize_status_value(cell.value)
        if cleaned != cell.value:
            cell.value = cleaned
            normalized += 1
    return normalized


def normalize_vf_cells(ws, header_row: int) -> int:
    headers = normalized_header_map(ws, header_row)
    project_col = headers.get("Project ID")
    vf_col = headers.get("VF or Non-VF")
    if project_col is None or vf_col is None:
        return 0

    normalized = 0
    for row in range(header_row + 1, ws.max_row + 1):
        if not normalize(ws.cell(row, project_col).value):
            continue
        cell = ws.cell(row, vf_col)
        cleaned = normalize_vf_value(cell.value)
        if cleaned != cell.value:
            cell.value = cleaned
            normalized += 1
    return normalized


def beneficiary_category_rows() -> list[tuple[str, str]]:
    return [("Beneficiary Category", "Indicator"), *BENEFICIARY_CATEGORIES]


def existing_beneficiary_category_rows(ws) -> list[tuple[str, str]]:
    rows: list[tuple[str, str]] = []
    for row in range(1, ws.max_row + 1):
        category = normalize(ws.cell(row, 1).value)
        indicator = normalize(ws.cell(row, 2).value)
        if category or indicator:
            rows.append((category, indicator))
    return rows


def update_beneficiary_categories_sheet(workbook) -> bool:
    expected_rows = beneficiary_category_rows()
    if BENEFICIARY_CATEGORIES_SHEET in workbook.sheetnames:
        ws = workbook[BENEFICIARY_CATEGORIES_SHEET]
    else:
        ws = workbook.create_sheet(BENEFICIARY_CATEGORIES_SHEET)

    if existing_beneficiary_category_rows(ws) == expected_rows:
        return False

    if ws.max_row:
        ws.delete_rows(1, ws.max_row)

    for row_idx, row_values in enumerate(expected_rows, start=1):
        for col_idx, value in enumerate(row_values, start=1):
            cell = ws.cell(row_idx, col_idx)
            cell.value = value
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if row_idx == 1:
                cell.font = Font(bold=True)

    ws.column_dimensions["A"].width = 36
    ws.column_dimensions["B"].width = 110
    ws.freeze_panes = "A2"
    return True


def category_is_mw(category: str) -> bool:
    return bool(MW_CATEGORY_RE.search(category))


def category_is_access_or_productive(category: str) -> bool:
    text = category.lower()
    if not text or category_is_mw(category):
        return False
    if any(term in text for term in EXCLUDE_CATEGORY_TERMS):
        return False
    return any(term in text for term in INCLUDE_CATEGORY_TERMS)


def has_numeric_value(value: object) -> bool:
    if value is None:
        return False
    if isinstance(value, (int, float)):
        return float(value) != 0
    text = str(value).strip()
    if not text:
        return False
    for match in NUMBER_RE.findall(text):
        try:
            if float(match.replace(",", "")) != 0:
                return True
        except ValueError:
            continue
    return False


def reported_number_columns(headers: dict[str, int]) -> list[int]:
    cols: list[int] = []
    for header, col in headers.items():
        header_text = header.lower()
        if header in {"Target", "End of Project Target"}:
            cols.append(col)
        elif "beneficiar" in header_text or "population" in header_text:
            cols.append(col)
    return sorted(set(cols))


def append_note(existing: object, note: str) -> str:
    text = normalize(existing)
    if note in text:
        return text
    return f"{text}\n{note}" if text else note


def fill_cells(ws, row: int, cols: Iterable[int], fill: PatternFill) -> None:
    for col in cols:
        if col is not None:
            ws.cell(row, col).fill = copy.copy(fill)


def add_cell_comments(ws, row: int, cols: Iterable[int], text: str) -> None:
    for col in cols:
        if col is not None:
            ws.cell(row, col).comment = Comment(text, "CO validation update")


def highlight_project_date_status_cells(ws, header_row: int) -> int:
    headers = normalized_header_map(ws, header_row)
    project_col = headers.get("Project ID")
    date_status_cols = [
        headers[header] for header in DATE_STATUS_HEADERS if header in headers
    ]
    if project_col is None or not date_status_cols:
        return 0

    highlighted = 0
    for row in range(header_row + 1, ws.max_row + 1):
        if normalize(ws.cell(row, project_col).value):
            fill_cells(ws, row, date_status_cols, ACTION_NEEDED_FILL)
            highlighted += 1
    return highlighted


def clear_project_info_on_output_rows(ws, header_row: int) -> int:
    headers = normalized_header_map(ws, header_row)
    project_col = headers.get("Project ID")
    project_info_cols = [
        headers[header]
        for header in PROJECT_INFO_HEADERS
        if header in headers
    ]
    if project_col is None or not project_info_cols:
        return 0

    cleared = 0
    for row in range(header_row + 1, ws.max_row + 1):
        if normalize(ws.cell(row, project_col).value):
            continue
        row_cleared = False
        for col in project_info_cols:
            cell = ws.cell(row, col)
            if normalize(cell.value) or cell.hyperlink:
                cell.value = None
                cell.hyperlink = None
                row_cleared = True
        if row_cleared:
            cleared += 1
    return cleared


def clear_cells(ws, row: int, cols: Iterable[int]) -> None:
    for col in cols:
        if col is not None:
            cell = ws.cell(row, col)
            cell.value = None
            cell.hyperlink = None


def iter_project_groups(ws, header_row: int, headers: dict[str, int]) -> Iterable[list[int]]:
    project_col = headers["Project ID"]
    groups: list[list[int]] = []
    current_group: list[int] = []

    for row in range(header_row + 1, ws.max_row + 1):
        row_values = [ws.cell(row, col).value for col in range(1, ws.max_column + 1)]
        if not any(value is not None and str(value).strip() for value in row_values):
            if current_group:
                groups.append(current_group)
                current_group = []
            continue

        if normalize(ws.cell(row, project_col).value):
            if current_group:
                groups.append(current_group)
            current_group = [row]
        elif current_group:
            current_group.append(row)

    if current_group:
        groups.append(current_group)

    return groups


def first_nonblank(ws, rows: list[int], col: int | None) -> object:
    if col is None:
        return None
    for row in rows:
        value = ws.cell(row, col).value
        if normalize(value):
            return value
    return None


def copy_row_style(ws, source_row: int, target_row: int) -> None:
    for col in range(1, ws.max_column + 1):
        source = ws.cell(source_row, col)
        target = ws.cell(target_row, col)
        if source.has_style:
            target._style = copy.copy(source._style)
        target.number_format = source.number_format
        target.alignment = copy.copy(source.alignment)
        target.font = copy.copy(source.font)
        target.border = copy.copy(source.border)


def add_or_mark_mw_rows(ws, header_row: int) -> tuple[int, int]:
    headers = normalized_header_map(ws, header_row)
    required = {"Project ID", "Beneficiary Category"}
    if not required.issubset(headers):
        return 0, 0

    target_col = headers.get("Target") or headers.get("End of Project Target")
    notes_col = headers.get("Notes") or headers.get("Comments") or headers.get("Comment")
    status_col = headers.get("Status")
    category_col = headers["Beneficiary Category"]
    indicator_col = headers.get("Indicator") or headers.get("Objective Indicator")
    baseline_col = headers.get("Baseline")
    reported_cols = reported_number_columns(headers)
    project_info_cols = [
        headers[header]
        for header in PROJECT_INFO_HEADERS
        if header in headers
    ]
    metadata_cols = [
        col
        for header in ("Notes", "Comments", "Comment", "Donor", "Donors", "Gender (% female)")
        if (col := headers.get(header)) is not None
    ]
    placeholder_blank_cols = sorted(set([*project_info_cols, *metadata_cols]))
    mw_issue_cols = [col for col in (baseline_col, target_col) if col is not None]

    groups_to_add: list[list[int]] = []
    mw_rows_to_mark: list[int] = []

    for group in iter_project_groups(ws, header_row, headers):
        access_rows_with_numbers = [
            row
            for row in group
            if category_is_access_or_productive(normalize(ws.cell(row, category_col).value))
            and any(has_numeric_value(ws.cell(row, col).value) for col in reported_cols)
        ]
        if not access_rows_with_numbers:
            continue

        mw_rows = [
            row
            for row in group
            if category_is_mw(normalize(ws.cell(row, category_col).value))
        ]
        mw_rows_with_value = [
            row
            for row in mw_rows
            if target_col is not None and has_numeric_value(ws.cell(row, target_col).value)
        ]
        if mw_rows_with_value:
            continue

        existing_auto_placeholder = False
        for row in mw_rows:
            note_text = normalize(ws.cell(row, notes_col).value) if notes_col else ""
            if AUTO_NOTE in note_text:
                existing_auto_placeholder = True
            mw_rows_to_mark.append(row)

        if not mw_rows and not existing_auto_placeholder:
            groups_to_add.append(group)

    marked_count = 0
    for row in sorted(set(mw_rows_to_mark)):
        fill_cells(ws, row, mw_issue_cols, ACTION_NEEDED_FILL)
        add_cell_comments(ws, row, mw_issue_cols, AUTO_NOTE)
        marked_count += 1

    added_count = 0
    for group in sorted(groups_to_add, key=lambda rows: rows[-1], reverse=True):
        insert_at = group[-1] + 1
        source_row = group[-1]
        ws.insert_rows(insert_at)
        copy_row_style(ws, source_row, insert_at)
        clear_cells(ws, insert_at, placeholder_blank_cols)

        ws.cell(insert_at, category_col).value = "Energy (MW added)"
        if indicator_col:
            ws.cell(insert_at, indicator_col).value = (
                "Installed capacity in MW for access/productive-use outputs"
            )
        if baseline_col:
            ws.cell(insert_at, baseline_col).value = 0
        if target_col:
            ws.cell(insert_at, target_col).value = None

        fill_cells(ws, insert_at, mw_issue_cols, ACTION_NEEDED_FILL)
        add_cell_comments(ws, insert_at, mw_issue_cols, AUTO_NOTE)
        added_count += 1

    return added_count, marked_count


def should_update_sheet(ws, sheet_name: str, all_project_sheets: bool) -> bool:
    if all_project_sheets:
        return True
    return ws.title == sheet_name


def update_workbook(
    path: Path, sheet_name: str, all_project_sheets: bool
) -> tuple[bool, list[SheetChange], openpyxl.Workbook]:
    workbook = openpyxl.load_workbook(path, data_only=True)
    changes: list[SheetChange] = []

    if update_beneficiary_categories_sheet(workbook):
        changes.append(
            SheetChange(
                workbook=path,
                sheet=BENEFICIARY_CATEGORIES_SHEET,
                beneficiary_categories_updated=True,
            )
        )

    for ws in workbook.worksheets:
        if not should_update_sheet(ws, sheet_name, all_project_sheets):
            continue
        header_row = find_header_row(ws)
        if header_row is None:
            continue

        change = SheetChange(workbook=path, sheet=ws.title)
        unmerge_sheet_cells(ws)
        change.columns_added = ensure_project_info_columns(ws, header_row)
        added, marked = add_or_mark_mw_rows(ws, header_row)
        change.placeholder_rows_added = added
        change.existing_mw_rows_marked = marked
        change.project_rows_highlighted = highlight_project_date_status_cells(
            ws, header_row
        )
        change.output_rows_project_info_cleared = clear_project_info_on_output_rows(
            ws, header_row
        )
        change.status_values_normalized = normalize_status_cells(ws, header_row)
        change.vf_values_normalized = normalize_vf_cells(ws, header_row)
        change.status_validation_added = ensure_status_data_validation(ws, header_row)
        if change.changed:
            changes.append(change)

    return bool(changes), changes, workbook


def output_workbook_path(path: Path, input_dir: Path, output_dir: Path) -> Path:
    relative = path.relative_to(input_dir)
    output_path = output_dir / relative
    output_path.parent.mkdir(parents=True, exist_ok=True)
    return output_path


def scrub_saved_project_info(output_path: Path, sheet_name: str, all_project_sheets: bool) -> int:
    workbook = openpyxl.load_workbook(output_path)
    cleared = 0
    for ws in workbook.worksheets:
        if not should_update_sheet(ws, sheet_name, all_project_sheets):
            continue
        header_row = find_header_row(ws)
        if header_row is None:
            continue
        cleared += clear_project_info_on_output_rows(ws, header_row)
    if cleared:
        workbook.save(output_path)
    return cleared


def workbook_paths(input_dir: Path) -> list[Path]:
    return sorted(
        path
        for path in input_dir.glob("*.xlsx")
        if path.is_file() and not path.name.startswith("~$")
    )


def main() -> int:
    args = parse_args()
    input_dir = args.input_dir
    if not input_dir.exists():
        raise SystemExit(f"Input directory does not exist: {input_dir}")

    mode = "APPLY" if args.apply else "DRY RUN"
    print(f"{mode}: scanning {input_dir}")

    all_changes: list[SheetChange] = []
    changed_workbooks = 0

    for path in workbook_paths(input_dir):
        changed, changes, workbook = update_workbook(
            path, args.sheet_name, args.all_project_sheets
        )
        if not changed:
            continue

        changed_workbooks += 1
        all_changes.extend(changes)

        if args.apply:
            output_path = output_workbook_path(path, input_dir, args.output_dir)
            workbook.save(output_path)
            scrub_saved_project_info(output_path, args.sheet_name, args.all_project_sheets)

    print(f"Workbooks with changes: {changed_workbooks}")
    print(f"Sheets with changes: {len(all_changes)}")
    print(f"Project-info columns inserted/moved: {sum(1 for c in all_changes if c.columns_added)} sheets")
    print(f"MW placeholder rows added: {sum(c.placeholder_rows_added for c in all_changes)}")
    print(f"Existing MW rows marked: {sum(c.existing_mw_rows_marked for c in all_changes)}")
    print(f"Project date/status rows highlighted: {sum(c.project_rows_highlighted for c in all_changes)}")
    print(
        "Output rows with repeated project info cleared: "
        f"{sum(c.output_rows_project_info_cleared for c in all_changes)}"
    )
    print(
        "Legacy status values normalized: "
        f"{sum(c.status_values_normalized for c in all_changes)}"
    )
    print(
        "VF/Non-VF labels normalized: "
        f"{sum(c.vf_values_normalized for c in all_changes)}"
    )
    print(
        "Status dropdown validations added: "
        f"{sum(1 for c in all_changes if c.status_validation_added)} sheets"
    )
    print(
        "Beneficiary category tabs updated: "
        f"{sum(1 for c in all_changes if c.beneficiary_categories_updated)}"
    )

    if all_changes:
        print("\nChanged sheets:")
        for change in all_changes:
            details = []
            if change.columns_added:
                details.append("project-info columns")
            if change.placeholder_rows_added:
                details.append(f"{change.placeholder_rows_added} MW placeholder row(s)")
            if change.existing_mw_rows_marked:
                details.append(f"{change.existing_mw_rows_marked} existing MW row(s) marked")
            if change.project_rows_highlighted:
                details.append(
                    f"{change.project_rows_highlighted} project date/status row(s) highlighted"
                )
            if change.output_rows_project_info_cleared:
                details.append(
                    f"{change.output_rows_project_info_cleared} output row(s) project info cleared"
                )
            if change.status_values_normalized:
                details.append(
                    f"{change.status_values_normalized} legacy status value(s) normalized"
                )
            if change.vf_values_normalized:
                details.append(
                    f"{change.vf_values_normalized} VF/Non-VF label(s) normalized"
                )
            if change.status_validation_added:
                details.append("status dropdown validation")
            if change.beneficiary_categories_updated:
                details.append("beneficiary category tab")
            print(f"- {change.workbook.name} :: {change.sheet}: {', '.join(details)}")

    if not args.apply:
        print("\nNo files were saved. Re-run with --apply to write changes.")
    else:
        print(f"\nUpdated workbooks exported to: {args.output_dir}")

    return 0


if __name__ == "__main__":
    raise SystemExit(main())
